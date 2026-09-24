"""
Batch runner: walks experiments/miplib_benchmark/, runs the 4 feasibility-pump
variants (via test_minlplib.py) on every instance, and appends each result to
experiments/benchmark_results.xlsx.

Resumable: instances already present in the target sheet are skipped, so this
can be interrupted and re-run safely.

Each instance runs in its own subprocess with a wall-clock timeout so one
huge/pathological instance can't stall the whole batch; timeouts and crashes
are logged as a marker row (Method="ALL") instead of stopping the run.

Usage:
    python experiments/run_all_benchmarks.py [--timeout SECONDS] [--limit N]
    python experiments/run_all_benchmarks.py --rerun-timeouts --timeout 3600
        Re-runs every instance currently marked "timeout"/"error" in the main
        sheet (with a longer budget) and files the results into the separate
        "Large Instances" sheet instead, removing their stale marker rows
        from the main sheet.
"""
from __future__ import annotations
import argparse
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from results_logger import (
    RESULTS_PATH, log_error, remove_instance_rows, format_duration,
    MAIN_SHEET, LARGE_SHEET, PERTURB_SHEET,
)
from instance_info import classify_type

BENCHMARK_DIR = Path(__file__).parent / "miplib_benchmark"
WORKER = Path(__file__).parent / "test_minlplib.py"
PERTURB_WORKER = Path(__file__).parent / "test_minlplib_perturb.py"
LOG_FILE = Path(__file__).parent / "run_all_benchmarks.log"

# .mps.gz and .mps duplicates exist for many instances; prefer the .gz copy.
EXT_PRIORITY = [".mps.gz", ".mps", ".lp.gz", ".lp"]


def discover_instances() -> list[Path]:
    groups: dict[str, dict[str, Path]] = {}
    for p in BENCHMARK_DIR.iterdir():
        if not p.is_file():
            continue
        for ext in EXT_PRIORITY:
            if p.name.endswith(ext):
                groups.setdefault(p.name[: -len(ext)], {})[ext] = p
                break

    chosen = []
    for base in sorted(groups):
        by_ext = groups[base]
        for ext in EXT_PRIORITY:
            if ext in by_ext:
                chosen.append(by_ext[ext])
                break
    return chosen


def base_name(p: Path) -> str:
    n = p.name
    for ext in EXT_PRIORITY:
        if n.endswith(ext):
            return n[: -len(ext)]
    return p.stem


def instances_in_sheet(sheet_name: str) -> set[str]:
    if not RESULTS_PATH.exists():
        return set()
    wb = load_workbook(RESULTS_PATH, read_only=True)
    if sheet_name not in wb.sheetnames:
        return set()
    ws = wb[sheet_name]
    names = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            names.add(row[0])
    return names


def instances_with_status(sheet_name: str, statuses: set[str]) -> set[str]:
    """Instance names in `sheet_name` whose Result column matches one of `statuses`."""
    if not RESULTS_PATH.exists():
        return set()
    wb = load_workbook(RESULTS_PATH, read_only=True)
    if sheet_name not in wb.sheetnames:
        return set()
    ws = wb[sheet_name]
    names = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        # columns: Instance, Type, Method, Iterations, Cuts Added, Perturbations, Result, Time, 비고
        if row and row[0] and row[6] in statuses:
            names.add(row[0])
    return names


def log(msg: str) -> None:
    line = f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {msg}"
    print(line, flush=True)
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(line + "\n")


def run_one(path: Path, timeout: float, sheet_name: str, index: int, total: int,
            worker: Path = WORKER) -> None:
    name = base_name(path)
    log(f"[{index}/{total}] {name} 시작...")
    t0 = time.perf_counter()
    inst_type = "?"
    try:
        inst_type = classify_type(path)
    except Exception:
        pass  # worker will still try to classify; fall back to "?" for marker rows only

    try:
        proc = subprocess.run(
            [sys.executable, str(worker), str(path), sheet_name, inst_type],
            timeout=timeout,
            capture_output=True,
            text=True,
        )
        elapsed = time.perf_counter() - t0
        if proc.returncode != 0:
            tail = (proc.stderr or "").strip().splitlines()
            note = tail[-1] if tail else f"exit code {proc.returncode}"
            log_error(name, inst_type, "error", elapsed, note[:200], sheet_name=sheet_name)
            log(f"[{index}/{total}] {name} 실패: {note[:200]}")
        else:
            log(f"[{index}/{total}] {name} 완료 ({format_duration(elapsed)})")
    except subprocess.TimeoutExpired:
        elapsed = time.perf_counter() - t0
        log_error(name, inst_type, "timeout", elapsed, f"{timeout}s 초과로 중단됨", sheet_name=sheet_name)
        log(f"[{index}/{total}] {name} 타임아웃 ({timeout}s)")
    except Exception as e:
        elapsed = time.perf_counter() - t0
        log_error(name, inst_type, "error", elapsed, str(e)[:200], sheet_name=sheet_name)
        log(f"[{index}/{total}] {name} 예외 발생: {e}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--timeout", type=float, default=600,
                         help="Per-instance wall-clock timeout in seconds (default: 600)")
    parser.add_argument("--limit", type=int, default=None,
                         help="Only process the first N remaining instances")
    parser.add_argument("--rerun-timeouts", action="store_true",
                         help="Re-run instances marked timeout/error in the main sheet, "
                              "with a longer budget, filing results into the Large Instances sheet")
    parser.add_argument("--resume-not-run", action="store_true",
                         help="Resume instances marked not_run in the Large Instances sheet "
                              "(left over from an interrupted --rerun-timeouts run)")
    parser.add_argument("--perturb", action="store_true",
                         help="Run the all-with-perturbation comparison (test_minlplib_perturb.py) "
                              "and log into the Perturbation Results sheet")
    args = parser.parse_args()

    if args.perturb and (args.rerun_timeouts or args.resume_not_run):
        parser.error("--perturb cannot be combined with --rerun-timeouts/--resume-not-run")

    instances = discover_instances()
    by_base = {base_name(p): p for p in instances}
    worker = WORKER

    if args.perturb:
        worker = PERTURB_WORKER
        # Instances tracked in the Large Instances sheet are the big ones that timed out
        # at 10 min; they are handled separately, so leave them out of this pass.
        done = instances_in_sheet(PERTURB_SHEET)
        large = instances_in_sheet(LARGE_SHEET)
        targets = [p for p in instances if base_name(p) not in done | large]
        if args.limit:
            targets = targets[: args.limit]
        target_sheet = PERTURB_SHEET
        log(f"[perturb] 총 {len(instances)}개 인스턴스 중 {len(done)}개 이미 완료, Large 시트 {len(large)}개 제외, "
            f"{len(targets)}개 실행 예정 (timeout={args.timeout}s)")
    elif args.resume_not_run:
        pending = instances_with_status(LARGE_SHEET, {"not_run"})
        targets = [by_base[n] for n in sorted(pending) if n in by_base]
        if args.limit:
            targets = targets[: args.limit]
        target_sheet = LARGE_SHEET
        removed = remove_instance_rows({base_name(p) for p in targets}, sheet_name=LARGE_SHEET)
        log(f"'{LARGE_SHEET}' 시트 not_run {len(pending)}개 중 {len(targets)}개 재개 (마커 {removed}개 제거), "
            f"timeout={args.timeout}s")
    elif args.rerun_timeouts:
        # Skip instances already re-run into the Large Instances sheet (resumable).
        stale = instances_with_status(MAIN_SHEET, {"timeout", "error"})
        already_retried = instances_in_sheet(LARGE_SHEET)
        stale -= already_retried
        targets = [by_base[n] for n in sorted(stale) if n in by_base]
        if args.limit:
            targets = targets[: args.limit]
        target_sheet = LARGE_SHEET
        # Only clear main-sheet marker rows for instances we're about to actually retry.
        removed = remove_instance_rows({base_name(p) for p in targets}, sheet_name=MAIN_SHEET)
        log(f"메인 시트 timeout/error {len(stale)}개 중 {len(targets)}개 재실행 (행 {removed}개 제거) → "
            f"'{LARGE_SHEET}' 시트에 timeout={args.timeout}s로 기록")
    else:
        done = instances_in_sheet(MAIN_SHEET)
        targets = [p for p in instances if base_name(p) not in done]
        if args.limit:
            targets = targets[: args.limit]
        target_sheet = MAIN_SHEET
        log(f"총 {len(instances)}개 인스턴스 중 {len(done)}개는 이미 완료됨, {len(targets)}개 실행 예정 "
            f"(timeout={args.timeout}s)")

    for i, path in enumerate(targets, start=1):
        run_one(path, args.timeout, target_sheet, i, len(targets), worker)

    log("=== 전체 배치 종료 ===")


if __name__ == "__main__":
    main()
