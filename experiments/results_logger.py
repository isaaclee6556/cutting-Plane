"""
Appends feasibility-pump benchmark results to a running Excel log
(experiments/benchmark_results.xlsx). Each run adds one row per method
for the instance just tested; the file is created with headers on first use.

Multiple sheets are supported (sheet_name param) so large/slow instances
can be collected separately from the main run.
"""
from __future__ import annotations
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font

RESULTS_PATH = Path(__file__).parent / "benchmark_results.xlsx"
MAIN_SHEET = "Benchmark Results"
LARGE_SHEET = "Large Instances"
PERTURB_SHEET = "Perturbation Results"
HEADERS = ["Instance", "Type", "Method", "Iterations", "Cuts Added", "Perturbations", "Result", "Time", "비고"]
COLUMN_WIDTHS = [22, 8, 16, 11, 11, 13, 14, 10, 45]


def format_duration(seconds: float) -> str:
    minutes, secs = divmod(seconds, 60)
    return f"{int(minutes)}m {secs:.2f}s"


def _remark(reason: str, perturbations: int) -> str:
    if reason == "feasible":
        return "사이클 발생, perturbation으로 회피" if perturbations > 0 else ""
    if reason == "cycle":
        return "사이클 발생 후 미해결 (perturbation 부족 또는 미사용)"
    if reason == "max_iter":
        return "최대 반복 횟수 도달, 미해결"
    if reason == "lp_infeasible":
        return "LP 완화 문제 자체가 infeasible"
    return ""


def _open_sheet(wb, sheet_name: str):
    if sheet_name in wb.sheetnames:
        return wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    return ws


def _load_or_create_workbook():
    if RESULTS_PATH.exists():
        return load_workbook(RESULTS_PATH)
    wb = Workbook()
    wb.remove(wb.active)  # default blank sheet; real sheets created via _open_sheet
    return wb


def _finalize_sheet(ws) -> None:
    last_col = get_column_letter(len(HEADERS))
    ws.auto_filter.ref = f"A1:{last_col}{ws.max_row}"
    for i, w in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def log_results(instance_name: str, instance_type: str, rows: list[dict],
                 sheet_name: str = MAIN_SHEET) -> None:
    """
    rows: list of dicts, each with keys
        method (str), iterations (int), cuts_added (int),
        perturbations (int), reason (str), time_sec (float)
    """
    wb = _load_or_create_workbook()
    ws = _open_sheet(wb, sheet_name)

    for r in rows:
        ws.append([
            instance_name,
            instance_type,
            r["method"],
            r["iterations"],
            r.get("cuts_added", 0),
            r.get("perturbations", 0),
            r["reason"],
            format_duration(r["time_sec"]),
            "; ".join(s for s in (_remark(r["reason"], r.get("perturbations", 0)), r.get("note", "")) if s),
        ])

    _finalize_sheet(ws)
    wb.save(RESULTS_PATH)
    print(f"\n[결과 저장] {RESULTS_PATH} ({sheet_name})")


def log_error(instance_name: str, instance_type: str, status: str, elapsed_sec: float,
              note: str, sheet_name: str = MAIN_SHEET) -> None:
    """
    Record a single marker row for an instance that could not be fully
    benchmarked (timeout, crash, license error, etc.), so the batch runner
    can skip it on resume instead of retrying it forever.
    """
    wb = _load_or_create_workbook()
    ws = _open_sheet(wb, sheet_name)

    ws.append([instance_name, instance_type, "ALL", "", "", "", status, format_duration(elapsed_sec), note])

    _finalize_sheet(ws)
    wb.save(RESULTS_PATH)


def remove_instance_rows(instance_names: set[str], sheet_name: str = MAIN_SHEET) -> int:
    """Delete all rows for the given instance names from a sheet. Returns count removed."""
    if not RESULTS_PATH.exists():
        return 0
    wb = load_workbook(RESULTS_PATH)
    if sheet_name not in wb.sheetnames:
        return 0
    ws = wb[sheet_name]

    rows_to_delete = [
        row[0].row for row in ws.iter_rows(min_row=2)
        if row[0].value in instance_names
    ]
    for row_idx in sorted(rows_to_delete, reverse=True):
        ws.delete_rows(row_idx)

    _finalize_sheet(ws)
    wb.save(RESULTS_PATH)
    return len(rows_to_delete)
